import csv
import io
from datetime import date, datetime

from django.db import transaction

from planner.models import Product, Sale

HEADER_ALIASES = {
    "date": "date",
    "day": "date",
    "sku": "sku",
    "product": "sku",
    "product_sku": "sku",
    "item": "sku",
    "quantity": "quantity",
    "qty": "quantity",
    "units": "quantity",
    "total_sales": "quantity",
    "sales": "quantity",
    "name": "name",
    "product_name": "name",
    "category": "category",
    "unit_cost": "unit_cost",
    "cost": "unit_cost",
    "lead_time_days": "lead_time_days",
    "lead_time": "lead_time_days",
    "on_hand": "on_hand",
    "stock": "on_hand",
}


def _normalize_header(value):
    key = str(value or "").strip().lower().replace(" ", "_")
    return HEADER_ALIASES.get(key)


def _parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%d %b %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Could not read the date '{value}'.")


def _rows_from_upload(uploaded):
    name = uploaded.name.lower()
    if name.endswith(".csv"):
        text = uploaded.read().decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise ValueError("The file needs a header row.")
        mapping = {_normalize_header(field): field for field in reader.fieldnames}
        mapping.pop(None, None)
        if "date" not in mapping or "quantity" not in mapping:
            raise ValueError("The file needs a date column and a units column (quantity or total_sales).")
        for raw in reader:
            yield {key: raw.get(source) for key, source in mapping.items()}
        return

    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        workbook = load_workbook(uploaded, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            header = next(iterator)
        except StopIteration as exc:
            raise ValueError("The workbook is empty.") from exc
        mapping = {}
        for index, cell in enumerate(header):
            key = _normalize_header(cell)
            if key:
                mapping[key] = index
        if "date" not in mapping or "quantity" not in mapping:
            raise ValueError("The file needs a date column and a units column (quantity or total_sales).")
        for raw in iterator:
            if raw is None or all(cell is None for cell in raw):
                continue
            yield {key: raw[index] if index < len(raw) else None for key, index in mapping.items()}
        return

    raise ValueError("Upload a .csv or .xlsx file.")


def import_sales(company, uploaded):
    created_products = 0
    written_sales = 0
    skipped = 0
    with transaction.atomic():
        for row in _rows_from_upload(uploaded):
            sku = str(row.get("sku") or "").strip() or "STORE"
            if row.get("date") in (None, "") or row.get("quantity") in (None, ""):
                skipped += 1
                continue
            try:
                sold_on = _parse_date(row["date"])
                quantity = int(round(float(row["quantity"])))
            except (TypeError, ValueError):
                skipped += 1
                continue
            if quantity < 0:
                skipped += 1
                continue
            product = company.products.filter(sku=sku).first()
            if product is None:
                product = Product.objects.create(
                    company=company,
                    sku=sku,
                    name=str(row.get("name") or ("Whole shop" if sku == "STORE" else sku)).strip(),
                    category=str(row.get("category") or "General").strip() or "General",
                    unit_cost=row.get("unit_cost") or 0,
                    lead_time_days=int(row.get("lead_time_days") or 7),
                    on_hand=int(row.get("on_hand") or 0),
                )
                created_products += 1
            Sale.objects.update_or_create(
                product=product,
                date=sold_on,
                defaults={"quantity": quantity},
            )
            written_sales += 1
    if written_sales == 0:
        raise ValueError("No sales rows were imported. Check the date and units columns.")
    return {"products": created_products, "sales": written_sales, "skipped": skipped}
